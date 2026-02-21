import openpyxl
from pathlib import Path
import psycopg2
import re

# Medical Terminology Mapping (Canonical Translations)
MEDICAL_TERMS = {
    "كشف": "Medical Consultation",
    "استشارة": "Medical Consultation",
    "متابعة": "Follow-up Visit",
    "جلسة": "Session",
    "طوارئ": "Emergency",
    "زيارة منزلية": "Home Visit",
    "ساعة": "Hour",
    "إقامة": "Inpatient Hospitalization",
    "تفتيت حصوات الكلى": "Lithotripsy",
    "أسنان": "Dental Care",
    "علاج طبيعي": "Physical Therapy",
    "نظارات": "Optical Service",
    "تحليل": "Laboratory Test",
    "أشعة": "Radiology",
    "تخدير": "Anesthesia",
    "إيواء": "Accommodation",
    "خدمات الرعاية بالعناية المركزة": "Intensive Care Services",
}

BLACKLIST = ["الخدمه", "التخصص", "التصنيف", "السعر", "العملية", "اسم الاجراء", "قيمة الاجراء", "price", "code", "Procedure"]

def clean_text(text):
    if not text: return ""
    text = str(text)
    text = re.sub(r'^[A-Z-]+[0-9.]+\s+', '', text)
    text = re.sub(r'[\.\s]+$', '', text)
    text = re.sub(r'\s+', ' ', text).strip()
    return text

def get_medical_translation(ar_name):
    ar_name_clean = clean_text(ar_name)
    if ar_name_clean == "الجلسة الأولى": return "1st Session"
    if ar_name_clean == "الجلسة الثانية": return "2nd Session"
    if ar_name_clean == "الجلسة الثالثة": return "3rd Session"
    
    for key, val in MEDICAL_TERMS.items():
        if key in ar_name_clean:
            if len(ar_name_clean) < 25: return val
    return ar_name_clean

def is_valid_service(name):
    if not name or len(name) < 2: return False
    if name in BLACKLIST: return False
    return True

def extract_from_venice(sheet):
    services = []
    current_specialization = "General"
    current_classification = "General"
    for row in sheet.iter_rows(values_only=True):
        if not any(row): continue
        col1 = str(row[0]) if row[0] is not None else ""
        col2 = str(row[1]) if row[1] is not None else ""
        
        # Venice Specific Header Detection
        if any(kw in col1 for kw in ["تفتيت", "أسنان", "كشف", "إقامة", "عمليات"]):
             current_specialization = clean_text(col1)
             if "عمليات" in col1: current_classification = "عمليات"
             continue
             
        if len(row) > 3 and row[3]:
            new_classif = clean_text(row[3])
            if new_classif not in BLACKLIST:
                current_classification = new_classif
        
        if col2 and not col1.startswith('Code') and col1 not in BLACKLIST:
            name_ar = clean_text(col2)
            if is_valid_service(name_ar):
                services.append({
                    "name_ar": name_ar,
                    "specialization": current_specialization,
                    "classification": current_classification
                })
    return services

def extract_from_dar_shifa_services(sheet):
    services = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if len(row) < 6: continue
        name_raw = clean_text(row[3])
        if is_valid_service(name_raw):
            services.append({
                "name_ar": name_raw,
                "specialization": clean_text(row[4]) or "General",
                "classification": clean_text(row[5]) or "General"
            })
    return services

def extract_from_dar_shifa_ops(sheet):
    services = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if len(row) < 15: continue
        name_raw = clean_text(row[5])
        if is_valid_service(name_raw):
            services.append({
                "name_ar": name_raw,
                "specialization": clean_text(row[12]) or "General",
                "classification": clean_text(row[14]) or "General"
            })
    return services

def extract_from_dental(sheet):
    services = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if len(row) < 3: continue
        name_raw = clean_text(row[0])
        if is_valid_service(name_raw):
            services.append({
                "name_ar": name_raw,
                "specialization": "Dental",
                "classification": clean_text(row[2]) or "Dental"
            })
    return services

def rebuild_dictionary():
    base_path = Path(r"d:\Backend\waadTbaSystem2026-main_final\waadTbaSystem2026-main\قائمة الخدمات المعتمدة")
    canonical_services = {}

    def add_to_canonical(extracted_list):
        for s in extracted_list:
            key = s["name_ar"].lower().strip()
            if key not in canonical_services:
                canonical_services[key] = s

    # 1. Venice
    path = base_path / "قائمة مفصلة اسعار مستشفى فينيسيا.xlsx"
    if path.exists():
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        add_to_canonical(extract_from_venice(wb.worksheets[0]))

    # 2. Dar Shifa Services
    path = base_path / "قائمة اسعار خدمات دار الشفاء مصنفة.xlsx"
    if path.exists():
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        add_to_canonical(extract_from_dar_shifa_services(wb.worksheets[0]))

    # 3. Dar Shifa Ops
    path = base_path / "قائمة اسعار عمليات دار الشفاء مصنفة.xlsx"
    if path.exists():
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        add_to_canonical(extract_from_dar_shifa_ops(wb.worksheets[0]))

    # 4. Dental
    path = base_path / "مركز دنتال لطب الاسنان.xlsx"
    if path.exists():
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        add_to_canonical(extract_from_dental(wb.worksheets[0]))

    try:
        conn = psycopg2.connect("host=localhost dbname=tba_waad_system user=postgres password=postgres port=5432")
        cur = conn.cursor()
        cur.execute("TRUNCATE TABLE ent_medical_services RESTART IDENTITY")
        print(f"Total Unique Services Loaded: {len(canonical_services)}")
        
        sorted_services = sorted(canonical_services.values(), key=lambda x: (x['classification'], x['specialization'], x['name_ar']))
        
        for idx, s in enumerate(sorted_services, start=1):
            seq_code = f"MED-{idx:04d}"
            name_ar = s["name_ar"]
            # Detect English names in Arabic field
            if re.match(r'^[a-zA-Z0-9[:space:]\/:,\.\-]+$', name_ar):
                name_en = name_ar
            else:
                name_en = get_medical_translation(name_ar)

            cur.execute("""
                INSERT INTO ent_medical_services (code, name_ar, name_en, category, sub_category)
                VALUES (%s, %s, %s, %s, %s)
            """, (seq_code, name_ar, name_en, s["specialization"], s["classification"]))
            
        conn.commit()
        print("Refined V3 Dictionary successfully populated.")
        cur.close()
        conn.close()
    except Exception as e:
        print(f"Database Error: {e}")

if __name__ == "__main__":
    rebuild_dictionary()
